#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Construit les donnees CHIFFREES de l'appli Stats a partir des exports Excel
(Export_Appli_AAAA-MM-JJ.xlsx) deposes dans le dossier source.

Rien de lisible n'est ecrit dans le depot : stats/data/ ne contient que des
fichiers AES-256-GCM (cle derivee du code d'acces par PBKDF2) :
    admin.json       tout le groupe (mot de passe admin)
    s_<code>.json    un magasin + ses vendeurs + les moyennes du groupe
                     (code d'acces du magasin)
    manifest.json    liste publique des magasins (aucun chiffre)

Les codes d'acces sont dans <source>/acces_stats.csv (cree au premier
lancement, a modifier a la main, a distribuer ; JAMAIS dans le depot). Les
codes sont pris tels quels (casse et symboles comptent). Apres modification,
relancer la construction : les appareils concernes ressaisissent le code.
La construction refuse un code ADMIN trop court ou identique au mot de passe
public de l'appli Suivi & Trophees.

Usage :
    py -3 build_stats.py
    py -3 build_stats.py --source "D:\\autre dossier"
"""

from __future__ import annotations

import argparse
import base64
import csv
import datetime as dt
import gzip
import hashlib
import json
import math
import os
import re
import secrets
import sys
from pathlib import Path

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent          # .../stats
APP = RACINE.parent                                      # racine de l'app web
KPIS_JSON = RACINE / "kpis.json"
SORTIE_DEF = RACINE / "data"
SOURCE_DEF = Path(r"C:\APP SANTE 2.0 stat")
PBKDF2_ITER = 250_000
ALPHABET = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"            # sans 0/O/1/I/L

# Magasins dont le libelle de l'export ne commence pas par leur numero.
ALIAS_DEFAUT = {
    "o2000 marsac": "37", "marsac": "37",
    "o2000 perigueux": "38", "o2000 px": "38",
    "go bergerac": "703", "go bretenoux": "570", "go tulle": "780",
}


def log(msg: str) -> None:
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


# ---------------------------------------------------------------------------
# Lecture / nettoyage
# ---------------------------------------------------------------------------
def num(v):
    """Nombre propre ou None (texte numerique, NaN, Infinity, #NAME?, vide)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).strip()
        if not s:
            return None
        try:
            x = float(s)
        except ValueError:
            return None
    if math.isnan(x) or math.isinf(x):
        return None
    return x


def arrondi(x):
    if x is None:
        return None
    return round(x, 5) if abs(x) < 10 else round(x, 2)


MOIS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def date_iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    m = re.search(r"(\d{1,2})\s+([A-Za-z]{3})\w*\s+(\d{4})", str(v or ""))
    if m and m.group(2).lower() in MOIS:
        return f"{int(m.group(3)):04d}-{MOIS[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", str(v or ""))
    return m.group(0) if m else None


def lignes(ws):
    it = ws.iter_rows(values_only=True)
    entete = [str(h).strip().strip("[]") if h is not None else "" for h in next(it)]
    return [dict(zip(entete, r)) for r in it if any(c is not None for c in r)]


def charger_json(chemin: Path, defaut):
    if chemin.exists():
        try:
            return json.loads(chemin.read_text(encoding="utf-8-sig"))
        except Exception as e:
            log(f"ATTENTION : {chemin.name} illisible ({e}), valeurs par defaut.")
    return defaut


# ---------------------------------------------------------------------------
# Referentiel des magasins (le meme que l'appli Suivi & Trophees)
# ---------------------------------------------------------------------------
def referentiel() -> dict:
    p = APP / "data" / "data.json"
    if not p.exists():
        return {}
    d = json.loads(p.read_text(encoding="utf-8"))
    return {s["code"]: s for s in d.get("stores", [])}


def code_depuis_libelle(libelle, alias) -> str | None:
    s = str(libelle or "").strip()
    if not s:
        return None
    m = re.match(r"(\d+)", s)
    if m:
        return str(int(m.group(1)))
    return alias.get(s.lower())


# ---------------------------------------------------------------------------
# Lecture d'un export
# ---------------------------------------------------------------------------
class Cat:
    def __init__(self, data: dict):
        self.data = data
        self.kpis = data["kpis"]
        self.keys = [k["key"] for k in self.kpis]
        self.by_key = {k["key"]: k for k in self.kpis}


def valeurs_ligne(r: dict, cat: Cat, o2000: bool) -> dict:
    """{kpi: (valeur, evolution)} pour une ligne de l'export."""
    out = {}
    for k in cat.kpis:
        col, evol = k["key"], k.get("evol")
        if o2000 and k.get("o2000"):
            col = k["o2000"]["col"]
            evol = k["o2000"].get("evol", evol)
        v = num(r.get(col))
        e = num(r.get(evol["key"])) if evol else None
        out[k["key"]] = (v, e)
    # Duo+ global : absent chez Optic 2000, reconstitue (libre/grille ponderes)
    if out["Tx_DuoPlus_Global"][0] is None:
        lib, gri = out["Tx_DuoPlus_Libre"][0], out["Tx_DuoPlus_Grille"][0]
        nl, ng = out["Nb_Ventes_Libre"][0] or 0, out["Nb_Ventes_Grille"][0] or 0
        if lib is not None and gri is not None and nl + ng > 0:
            out["Tx_DuoPlus_Global"] = ((lib * nl + gri * ng) / (nl + ng), None)
        elif lib is not None and nl > 0:
            out["Tx_DuoPlus_Global"] = (lib, None)
    return out


def lire_export(chemin: Path, cat: Cat, alias: dict) -> dict:
    wb = load_workbook(chemin, data_only=True)
    snap = {"fichier": chemin.name, "debut": None, "fin": None, "compta_fin": None,
            "groupe": None, "groupe_o2000": None, "magasins": {}, "compta": {}, "compta_groupe": None,
            "hors_magasin": []}

    def feuille(nom, o2000):
        if nom not in wb.sheetnames:
            log(f"  feuille '{nom}' absente de {chemin.name}")
            return
        for r in lignes(wb[nom]):
            niveau = str(r.get("Niveau") or "")
            snap["debut"] = snap["debut"] or date_iso(r.get("Date_Debut"))
            if not o2000:
                snap["fin"] = snap["fin"] or date_iso(r.get("Date_Fin"))
            vals = valeurs_ligne(r, cat, o2000)
            if niveau == "Groupe":
                snap["groupe"] = vals
            elif niveau.startswith("Groupe"):
                snap["groupe_o2000"] = vals
            elif niveau == "Magasin":
                code = code_depuis_libelle(r.get("Magasin"), alias)
                if not code:
                    log(f"  magasin sans code ignore : {r.get('Magasin')!r}")
                    continue
                m = snap["magasins"].setdefault(code, {"label": r.get("Magasin"), "o2000": o2000,
                                                       "vals": None, "vendeurs": []})
                m["vals"], m["o2000"] = vals, o2000
            elif niveau == "Vendeur":
                code = code_depuis_libelle(r.get("Magasin"), alias)
                if not code:
                    continue
                m = snap["magasins"].setdefault(code, {"label": r.get("Magasin"), "o2000": o2000,
                                                       "vals": None, "vendeurs": []})
                m["vendeurs"].append({"nom": str(r.get("Vendeur") or "").strip(), "vals": vals})

    feuille("Vendeurs", False)
    feuille("Optic2000", True)

    if "CA Compta" in wb.sheetnames:
        for r in lignes(wb["CA Compta"]):
            niveau = str(r.get("Niveau") or "")
            snap["compta_fin"] = snap["compta_fin"] or date_iso(r.get("Date_Fin"))
            rec = {"n2": num(r.get("CA_N2")), "n1": num(r.get("CA_N1")), "n": num(r.get("CA_N")),
                   "evol": num(r.get("CA_Evol")), "comp": num(r.get("CA_Comparable")),
                   "comp_evol": num(r.get("CA_Comparable_Evol"))}
            if niveau == "Groupe":
                snap["compta_groupe"] = rec
            elif niveau == "Magasin":
                code = code_depuis_libelle(r.get("Magasin"), alias) or alias.get(str(r.get("Magasin_CA") or "").lower())
                if code:
                    snap["compta"][code] = rec
                else:
                    log(f"  CA compta sans code ignore : {r.get('Magasin_CA')!r}")
            else:
                code = alias.get(str(r.get("Magasin_CA") or "").lower())
                if code:
                    snap["compta"].setdefault(code, rec)
                snap["hors_magasin"].append({"nom": r.get("Magasin_CA"), **rec})
    if not snap["fin"]:
        m = re.search(r"(\d{4}-\d{2}-\d{2})", chemin.name)
        snap["fin"] = m.group(1) if m else None
    return snap


# ---------------------------------------------------------------------------
# Score de sante (reference = instance superieure)
# ---------------------------------------------------------------------------
def borne(x, cfg):
    return max(cfg["indexMin"], min(cfg["indexMax"], x))


def zone_pour(score, zones):
    for z in zones:
        if score < z["max"]:
            return z["key"]
    return zones[-1]["key"]


def calcul_sante(vals: dict, ref: dict, cat: Cat) -> dict | None:
    if not vals or not ref:
        return None
    cfg = cat.data["health"]
    piliers, idx = {}, {}
    for p in cat.data["pillars"]:
        liste = []
        if p["mode"] == "growth":
            e, re_ = vals["CA"][1], ref["CA"][1]
            if e is not None and re_ is not None:
                i = borne(100 + cfg["growthSlope"] * (e - re_), cfg)
                idx["CA"] = i
                liste.append(i)
        else:
            for key in p["kpis"]:
                v, r = vals.get(key, (None, None))[0], ref.get(key, (None, None))[0]
                if v is None or r is None or r <= 0:
                    continue
                down = cat.by_key[key]["polarity"] == "down"
                if down and v <= 0:
                    i = cfg["indexMax"]
                else:
                    i = borne(100 * ((r / v) if down else (v / r)), cfg)
                idx[key] = i
                liste.append(i)
        if liste:
            piliers[p["key"]] = sum(liste) / len(liste)
    if len(piliers) < cfg["minPillars"]:
        return None
    poids = {p["key"]: p.get("weight", 1) for p in cat.data["pillars"]}
    tot = sum(poids[k] for k in piliers)
    score = sum(piliers[k] * poids[k] for k in piliers) / tot
    return {"score": round(score, 1), "zone": zone_pour(score, cat.data["zones"]),
            "piliers": {k: round(v, 1) for k, v in piliers.items()},
            "idx": {k: round(v, 1) for k, v in idx.items()}}


def vendeur_actif(vals: dict, nom: str, cfg: dict) -> bool:
    if nom in cfg.get("excludeSellers", []):
        return False
    ventes = (vals.get("Nb_Total") or (None, None))[0] or 0
    ca = (vals.get("CA") or (None, None))[0] or 0
    return ventes >= cfg["minSellerSales"] and ca >= cfg["minSellerCA"]


# ---------------------------------------------------------------------------
# Construction des enregistrements
# ---------------------------------------------------------------------------
def vers_listes(vals: dict | None, cat: Cat):
    if not vals:
        return None, None
    return ([arrondi(vals[k][0]) for k in cat.keys], [arrondi(vals[k][1]) for k in cat.keys])


def construire(snap: dict, cat: Cat, roster: dict, params: dict) -> dict:
    cfg = cat.data["health"]
    groupe = snap["groupe"]
    # Optic 2000 n'a pas les memes definitions de taux (Duo+, etc.) : ses
    # magasins se comparent au "Groupe Optic 2000", pas au groupe principal.
    groupe_o2000 = snap["groupe_o2000"] or groupe
    magasins = []
    for code in sorted(set(roster) | set(snap["magasins"]) | set(snap["compta"]), key=lambda c: (len(c), c)):
        info = roster.get(code, {})
        m = snap["magasins"].get(code)
        compta = snap["compta"].get(code)
        if not m and not compta:
            continue
        vals = m["vals"] if m else None
        kind = "ca" if not vals else ("o2000" if m["o2000"] else "full")
        ref = groupe_o2000 if kind == "o2000" else groupe
        v, e = vers_listes(vals, cat)
        sante = calcul_sante(vals, ref, cat) if vals else None
        ent = {"code": code, "nom": info.get("name") or (m or {}).get("label") or code,
               "enseigne": info.get("enseigne", ""), "ville": info.get("ville", ""),
               "kind": kind, "v": v, "e": e, "sante": sante,
               "compta": ({k: arrondi(x) for k, x in compta.items()} if compta else None),
               "vendeurs": []}
        if m:
            for s in m["vendeurs"]:
                sv, se = vers_listes(s["vals"], cat)
                actif = vendeur_actif(s["vals"], s["nom"], cfg)
                sh = calcul_sante(s["vals"], vals, cat) if (actif and vals) else None
                ent["vendeurs"].append({"nom": s["nom"], "v": sv, "e": se, "actif": actif, "sante": sh,
                                        "masque": s["nom"] in cfg.get("excludeSellers", [])})
            # rang des vendeurs actifs avec score, au sein du magasin
            classes = sorted([s for s in ent["vendeurs"] if s["sante"]], key=lambda s: -s["sante"]["score"])
            for i, s in enumerate(classes, 1):
                s["rang"] = [i, len(classes)]
            ent["vendeurs"].sort(key=lambda s: (s["masque"], not s["actif"], -(s["v"][0] or 0)))
        magasins.append(ent)
    # classement au sein de chaque reseau (full / o2000)
    for kind in ("full", "o2000"):
        classes = sorted([m for m in magasins if m["sante"] and m["kind"] == kind],
                         key=lambda m: -m["sante"]["score"])
        for i, m in enumerate(classes, 1):
            m["rang"] = [i, len(classes)]
    return {"magasins": magasins, "groupe": groupe, "groupe_o2000": groupe_o2000}


def distribution(magasins: list, cat: Cat) -> dict:
    """Valeurs triees des magasins du reseau principal, par indicateur (pour
    situer un magasin entre le min et le max, sans reveler qui est qui)."""
    dist = {}
    for i, key in enumerate(cat.keys):
        xs = sorted(m["v"][i] for m in magasins
                    if m["kind"] == "full" and m["v"] and m["v"][i] is not None)
        if len(xs) >= 3:
            dist[key] = xs
    return dist


# ---------------------------------------------------------------------------
# Chiffrement
# ---------------------------------------------------------------------------
def normaliser_code(s: str) -> str:
    # Le code est pris TEL QUEL (casse et symboles respectes, espaces de
    # debut/fin retires) : c'est ce qui garde toute sa force a un mot de passe
    # choisi. Meme regle cote navigateur (stats/js/crypto.js).
    return str(s).strip()


def b64(b: bytes) -> str:
    return base64.b64encode(b).decode("ascii")


def sel_fichier(sel_secret: str, nom: str) -> bytes:
    """Sel STABLE par fichier : la cle derivee reste la meme d'une publication
    a l'autre, ce qui permet a un appareil de la memoriser (le code n'est
    demande qu'a la premiere connexion). Le sel n'est pas un secret ; l'IV,
    lui, est aleatoire a chaque ecriture."""
    return hashlib.sha256(f"{sel_secret}|{nom}".encode("utf-8")).digest()[:16]


def chiffrer(obj, code: str, sel: bytes) -> dict:
    brut = gzip.compress(json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), 9)
    iv = os.urandom(12)
    cle = hashlib.pbkdf2_hmac("sha256", normaliser_code(code).encode("utf-8"), sel, PBKDF2_ITER, dklen=32)
    ct = AESGCM(cle).encrypt(iv, brut, None)
    return {"v": 1, "kdf": "PBKDF2-SHA256", "iter": PBKDF2_ITER, "salt": b64(sel), "iv": b64(iv), "ct": b64(ct), "gz": True}


def nouveau_code(n=12) -> str:
    s = "".join(secrets.choice(ALPHABET) for _ in range(n))
    return "-".join(s[i:i + 4] for i in range(0, n, 4))


def lire_csv_acces(chemin: Path):
    """acces_stats.csv (Acces;Code d'acces) -> (code admin, {code magasin: code}).
    Le libelle d'un magasin finit par son numero entre parentheses, ex.
    'PERIGUEUX (1)'. Tolere un fichier enregistre par Excel (; ou , / ANSI)."""
    texte = None
    for enc in ("utf-8-sig", "cp1252"):
        try:
            texte = chemin.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    lignes = texte.splitlines()
    delim = ";" if lignes and ";" in lignes[0] else ","
    admin, magasins = None, {}
    for r in list(csv.reader(lignes, delimiter=delim))[1:]:
        if len(r) < 2 or not r[1].strip():
            continue
        lib, code = r[0].strip(), r[1].strip()
        if lib.upper() == "ADMIN":
            admin = code
        else:
            m = re.search(r"\((\d+)\)\s*$", lib)
            if m:
                magasins[m.group(1)] = code
    return admin, magasins


# Empreintes (SHA-256) de mots de passe devenus publics : l'ancien mot de passe
# admin de l'appli Suivi & Trophees etait ecrit en clair dans son code, donc
# dans l'historique du depot GitHub public. Il ne doit plus jamais servir.
MOTS_DE_PASSE_BRULES = {
    "c823cc0c7a4a9538c861c7e3b1562ba99faf1d0e027d7ae1f0d2b7b255a6bb3c",
}


def est_brule(code: str) -> bool:
    return hashlib.sha256(code.encode("utf-8")).hexdigest() in MOTS_DE_PASSE_BRULES


def verifier_acces(acces: dict, noms: dict) -> tuple[list, list]:
    erreurs, avertissements = [], []
    admin = acces["admin"]
    if est_brule(admin):
        erreurs.append("Le code ADMIN est l'ancien mot de passe de l'appli Suivi & Trophees : il est lisible par "
                       "tout le monde dans l'historique public de GitHub, il ne protegerait rien. "
                       "Choisissez-en un nouveau (au moins 10 caracteres) : il servira aux deux applis.")
    elif len(admin) < 8:
        erreurs.append(f"Le code ADMIN est trop court ({len(admin)} caracteres) : les fichiers chiffres sont "
                       "publics, un code court se retrouve par force brute. Visez 12 caracteres ou plus.")
    elif len(admin) < 10:
        avertissements.append(f"Code ADMIN assez court ({len(admin)} caracteres) : 12 ou plus est plus sur.")
    codes = acces["magasins"]
    partages = {}
    for c, v in codes.items():
        partages.setdefault(v, []).append(c)
        if v == admin:
            erreurs.append(f"Le code du magasin {noms.get(c, c)} est identique au code ADMIN.")
        if est_brule(v):
            erreurs.append(f"Le code du magasin {noms.get(c, c)} est un ancien mot de passe devenu public.")
        if len(v) < 6:
            avertissements.append(f"Code trop court pour {noms.get(c, c)} ({len(v)} caracteres).")
    groupes = [cs for cs in partages.values() if len(cs) > 1]
    if groupes:
        n = max(len(g) for g in groupes)
        avertissements.append(f"{n} magasins partagent le meme code : l'appli empeche un magasin d'aller voir un "
                              "autre (pas de bouton, appareil verrouille), mais ce n'est pas une barriere technique : "
                              "un code par magasin en serait une.")
    return erreurs, avertissements


def charger_acces(source: Path, codes_magasins: list) -> dict:
    """Source de verite des codes : acces_stats.csv (a modifier a la main).
    acces_stats.json ne garde que le 'sel' (parametre technique stable)."""
    chemin_json = source / "acces_stats.json"
    chemin_csv = source / "acces_stats.csv"
    meta = charger_json(chemin_json, {})
    acces = {"admin": None, "magasins": {}}
    if chemin_csv.exists():
        acces["admin"], acces["magasins"] = lire_csv_acces(chemin_csv)
    else:                               # reprise d'un ancien acces_stats.json
        acces["admin"] = meta.get("admin")
        acces["magasins"] = dict(meta.get("magasins", {}))
    sel = meta.get("sel") or secrets.token_hex(16)
    ajoutes = []
    if not acces["admin"]:
        acces["admin"] = nouveau_code(); ajoutes.append("ADMIN")
    for c in codes_magasins:
        if not acces["magasins"].get(c):
            acces["magasins"][c] = nouveau_code(); ajoutes.append(c)
    acces["sel"] = sel
    if ajoutes:
        log(f"Codes generes pour : {', '.join(ajoutes)}")
    return acces


def ecrire_acces(source: Path, acces: dict, magasins: list) -> None:
    lignes_csv = ["Acces;Code d'acces", f"ADMIN;{acces['admin']}"] + [
        f"{m['nom']} ({m['code']});{acces['magasins'][m['code']]}" for m in magasins]
    (source / "acces_stats.csv").write_text("\n".join(lignes_csv), encoding="utf-8-sig")
    (source / "acces_stats.json").write_text(json.dumps(
        {"sel": acces["sel"], "_note": "Parametre technique : ne pas modifier. Les codes se changent dans acces_stats.csv."},
        indent=1), encoding="utf-8")


# ---------------------------------------------------------------------------
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Construit les donnees chiffrees de l'appli Stats")
    ap.add_argument("--source", type=Path, default=SOURCE_DEF, help="dossier des exports Excel")
    ap.add_argument("--sortie", type=Path, default=SORTIE_DEF)
    ap.add_argument("--debug-clair", type=Path, help="ecrit aussi le JSON admin NON chiffre ici (test local uniquement)")
    args = ap.parse_args(argv)

    cat = Cat(json.loads(KPIS_JSON.read_text(encoding="utf-8")))
    params = charger_json(args.source / "parametres_stats.json", {})
    alias = {**ALIAS_DEFAUT, **{k.lower(): str(v) for k, v in params.get("alias", {}).items()}}
    roster = referentiel()

    fichiers = sorted(args.source.glob("Export_Appli_*.xlsx"))
    if not fichiers:
        log(f"Aucun Export_Appli_*.xlsx dans {args.source}")
        return 2

    snaps = []
    for f in fichiers:
        log(f"Lecture {f.name}")
        try:
            s = lire_export(f, cat, alias)
        except PermissionError:
            log(f"{f.name} est ouvert dans Excel : fermez-le et relancez.")
            return 1
        if not s["groupe"]:
            log(f"  {f.name} : ligne Groupe absente, ignore.")
            continue
        snaps.append(s)
    snaps.sort(key=lambda s: s["fin"] or "")
    courant = snaps[-1]
    log(f"Periode courante : {courant['debut']} -> {courant['fin']} (CA compta jusqu'au {courant['compta_fin']})")

    built = construire(courant, cat, roster, params)
    magasins = built["magasins"]

    # historique : une courbe par magasin / vendeur sur les exports successifs
    snap_meta = [{"id": s["fin"], "debut": s["debut"], "fin": s["fin"], "compta_fin": s["compta_fin"],
                  "fichier": s["fichier"]} for s in snaps]
    if len(snaps) > 1:
        for s in snaps[:-1]:
            anc = construire(s, cat, roster, params)
            par_code = {m["code"]: m for m in anc["magasins"]}
            for m in magasins:
                a = par_code.get(m["code"])
                if not a:
                    continue
                m.setdefault("hist", []).append({"s": s["fin"], "ca": (a["v"] or [None])[0],
                                                 "score": a["sante"]["score"] if a["sante"] else None})
                for sv in m["vendeurs"]:
                    av = next((x for x in a["vendeurs"] if x["nom"] == sv["nom"]), None)
                    if av:
                        sv.setdefault("hist", []).append({"s": s["fin"], "ca": av["v"][0],
                                                          "score": av["sante"]["score"] if av["sante"] else None})
    for m in magasins:
        m.setdefault("hist", []).append({"s": courant["fin"], "ca": (m["v"] or [None])[0],
                                         "score": m["sante"]["score"] if m["sante"] else None})
        for sv in m["vendeurs"]:
            sv.setdefault("hist", []).append({"s": courant["fin"], "ca": sv["v"][0],
                                              "score": sv["sante"]["score"] if sv["sante"] else None})

    gv, ge = vers_listes(courant["groupe"], cat)
    g2v, g2e = vers_listes(built["groupe_o2000"], cat)
    commun = {
        "v": 1, "genere": dt.datetime.now().isoformat(timespec="seconds"),
        "snapshots": snap_meta, "courant": courant["fin"], "periode": {"debut": courant["debut"], "fin": courant["fin"],
                                                                         "compta_fin": courant["compta_fin"]},
        "kpis": cat.keys,
        "groupe": {"v": gv, "e": ge, "compta": courant["compta_groupe"],
                   "nb_magasins": sum(1 for m in magasins if m["kind"] == "full")},
        "groupe_o2000": {"v": g2v, "e": g2e, "nb_magasins": sum(1 for m in magasins if m["kind"] == "o2000")},
        "dist": distribution(magasins, cat),
    }

    codes = [m["code"] for m in magasins]
    acces = charger_acces(args.source, codes)
    erreurs, avertissements = verifier_acces(acces, {m["code"]: m["nom"] for m in magasins})
    for a in avertissements:
        log(f"ATTENTION : {a}")
    if erreurs:
        for e in erreurs:
            log(f"ERREUR : {e}")
        log(f"Rien n'a ete modifie. Corrigez {args.source / 'acces_stats.csv'} puis relancez.")
        return 3

    sortie = args.sortie
    sortie.mkdir(parents=True, exist_ok=True)
    ecrits = set()

    admin = {**commun, "role": "admin", "magasins": magasins, "hors_magasin": courant["hors_magasin"]}
    (sortie / "admin.json").write_text(
        json.dumps(chiffrer(admin, acces["admin"], sel_fichier(acces["sel"], "admin.json"))), encoding="utf-8")
    ecrits.add("admin.json")
    if args.debug_clair:
        args.debug_clair.write_text(json.dumps(admin, ensure_ascii=False), encoding="utf-8")

    for m in magasins:
        payload = {**commun, "role": "magasin", "magasins": [m]}
        nom = f"s_{m['code']}.json"
        (sortie / nom).write_text(
            json.dumps(chiffrer(payload, acces["magasins"][m["code"]], sel_fichier(acces["sel"], nom))), encoding="utf-8")
        ecrits.add(nom)

    manifest ={"v": 1, "genere": commun["genere"],
                "magasins": [{"code": m["code"], "nom": m["nom"], "ville": m["ville"], "enseigne": m["enseigne"],
                              "kind": m["kind"]} for m in magasins]}
    (sortie / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    ecrits.add("manifest.json")
    for f in sortie.glob("*.json"):
        if f.name not in ecrits:
            f.unlink()

    # codes d'acces : acces_stats.csv a cote des exports (hors depot)
    ecrire_acces(args.source, acces, magasins)

    taille = sum((sortie / n).stat().st_size for n in ecrits)
    n_vend = sum(len(m["vendeurs"]) for m in magasins)
    log(f"{len(magasins)} magasins, {n_vend} vendeurs -> {len(ecrits)} fichiers dans {sortie} ({taille // 1024} Ko)")
    scores = [m["sante"]["score"] for m in magasins if m["sante"]]
    if scores:
        log(f"Score de sante des magasins : min {min(scores)}, moyenne {sum(scores)/len(scores):.1f}, max {max(scores)}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nInterrompu.")
        sys.exit(130)
