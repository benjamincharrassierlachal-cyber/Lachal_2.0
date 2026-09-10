#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Exporte "Base suivi magasins.xlsx" (dossier Suivi Groupe) en data/data.json,
le fichier que lit l'appli web (index.html + js/).

Ne modifie jamais les fichiers sources : lecture seule, comme compiler_base.py.

Usage :
    py -3 exporter_donnees.py
    py -3 exporter_donnees.py --source "D:\\ailleurs\\Suivi Groupe"
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent           # racine de l'app
CHEMINS = RACINE / "scripts" / "chemins.json"              # override facultatif
SORTIE = RACINE / "data" / "data.json"


def log(msg: str) -> None:
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


def trouver_suivi_groupe() -> Path | None:
    """Retrouve le dossier 'Suivi Groupe' (celui qui contient le xlsx compile)."""
    if CHEMINS.exists():
        try:
            force = json.loads(CHEMINS.read_text(encoding="utf-8"))
            if force.get("SuiviGroupe"):
                p = Path(force["SuiviGroupe"])
                if (p / "Base suivi magasins.xlsx").exists():
                    return p
        except Exception:
            pass
    candidats = [
        RACINE.parent / "APP SANTE 2.0" / "Suivi Groupe",
        RACINE.parent / "APP SANTE 2.0" / "Suivi groupe",
        Path.home() / "Desktop" / "APP SANTE 2.0" / "Suivi Groupe",
        Path.home() / "Bureau" / "APP SANTE 2.0" / "Suivi Groupe",
    ]
    for base in (RACINE, RACINE.parent, RACINE.parent.parent,
                 Path.home() / "Desktop", Path.home() / "Bureau"):
        try:
            if not base.is_dir():
                continue
            for enfant in base.iterdir():
                if enfant.is_dir() and "suivi" in enfant.name.lower():
                    candidats.append(enfant)
                if enfant.is_dir() and "app sante" in enfant.name.lower():
                    for petit in enfant.iterdir():
                        if petit.is_dir() and "suivi" in petit.name.lower():
                            candidats.append(petit)
        except OSError:
            continue
    for c in candidats:
        if (c / "Base suivi magasins.xlsx").exists():
            return c
    return None


def en_iso(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10] if isinstance(v, dt.datetime) else v.isoformat()
    return None


def nombre(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    return None


def exporter(dossier: Path, sortie: Path) -> int:
    from openpyxl import load_workbook

    xlsx = dossier / "Base suivi magasins.xlsx"
    wb = load_workbook(xlsx, data_only=True, read_only=True)

    # --- Referentiel : liste des magasins ---
    ws_ref = wb["Referentiel"]
    entete_ref, stores = None, []
    for r in ws_ref.iter_rows(values_only=True):
        if entete_ref is None:
            entete_ref = list(r)
            continue
        if not r or not r[1]:
            continue
        col = dict(zip(entete_ref, r))
        stores.append({
            "code": str(col.get("Code") or "").strip(),
            "name": str(col.get("Magasin") or "").strip(),
            "enseigne": str(col.get("Enseigne") or "").strip(),
            "ville": str(col.get("Ville") or "").strip(),
        })

    # --- Objectifs ---
    ws_obj = wb["Objectifs"] if "Objectifs" in wb.sheetnames else None
    objectives = {}
    if ws_obj is not None:
        entete_obj = None
        for r in ws_obj.iter_rows(values_only=True):
            if entete_obj is None:
                entete_obj = list(r)
                continue
            if not r or not r[0]:
                continue
            col = dict(zip(entete_obj, r))
            code = str(col.get("Code") or "").strip()
            if not code:
                continue

            def oui(v):
                return str(v or "").strip().lower() == "oui"

            objectives[code] = {
                "avis_semaine": nombre(col.get("Objectif avis Google / semaine")),
                "note_cible": nombre(col.get("Note Google cible")),
                "examens_semaine": nombre(col.get("Objectif examens de vue / semaine (Lyleoo)")),
                "impressions_semaine": nombre(col.get("Objectif impressions 3D / semaine")),
                "show_avis": oui(col.get("Afficher avis Google")),
                "show_examens": oui(col.get("Afficher examens de vue")),
                "show_impressions": oui(col.get("Afficher impressions 3D")),
                "commentaire": str(col.get("Commentaire") or "").strip(),
            }

    # --- Base : historique hebdomadaire ---
    ws_base = wb["Base"]
    entete_base, brutes = None, []
    for r in ws_base.iter_rows(values_only=True):
        if entete_base is None:
            entete_base = list(r)
            continue
        if not r or not r[3]:
            continue
        brutes.append(dict(zip(entete_base, r)))

    weeks: dict[str, list] = {}
    for l in brutes:
        code = str(l.get("Code") or "").strip()
        if not code:
            continue
        faces, branches, goodies = (nombre(l.get("Faces")),
                                     nombre(l.get("Branches")),
                                     nombre(l.get("Goodies")))
        impressions = (None if faces is None and branches is None and goodies is None
                       else (faces or 0) + (branches or 0) + (goodies or 0))

        # "Demandes" issu d'un releve Lyleoo MENSUEL (historique d'avant le
        # passage en hebdomadaire) n'est pas comparable a un objectif
        # hebdomadaire : on l'exclut plutot que d'afficher un pourcentage
        # trompeur (ex. un mois entier compare a l'objectif d'une semaine).
        granularite = str(l.get("Granularite") or "")
        demandes = nombre(l.get("Demandes Lyleoo"))
        if demandes is not None and "mensuel" in granularite and "hebdomadaire" not in granularite:
            demandes = None

        semaine = {
            "semaine": l.get("Semaine"),
            "debut": en_iso(l.get("Debut semaine")),
            "fin": en_iso(l.get("Fin semaine")),
            "avis_nouveaux": nombre(l.get("Avis nouveaux")),
            "note": nombre(l.get("Note Google")),
            "avis_cumules": nombre(l.get("Avis cumules")),
            "faces": faces, "branches": branches, "goodies": goodies,
            "impressions": impressions,
            "demandes": demandes,
            "soumis": nombre(l.get("Dossiers soumis")),
            "valides": nombre(l.get("Dossiers valides")),
            "sources": l.get("Sources") or "",
        }
        weeks.setdefault(code, []).append(semaine)

    for code in weeks:
        weeks[code].sort(key=lambda s: s["semaine"] or "")

    data = {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "stores": sorted(stores, key=lambda s: s["name"]),
        "objectives": objectives,
        "weeks": weeks,
    }

    sortie.parent.mkdir(parents=True, exist_ok=True)
    sortie.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return len(brutes)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="Exporte la base magasins vers data/data.json")
    p.add_argument("--source", type=Path, help="dossier 'Suivi Groupe' (auto-detecte sinon)")
    p.add_argument("--sortie", type=Path, default=SORTIE)
    args = p.parse_args(argv)

    dossier = args.source or trouver_suivi_groupe()
    if not dossier or not (dossier / "Base suivi magasins.xlsx").exists():
        log("Dossier 'Suivi Groupe' introuvable (Base suivi magasins.xlsx absent). "
            "Creez scripts/chemins.json avec {\"SuiviGroupe\": \"chemin complet\"}.")
        return 2

    log(f"Source : {dossier}")
    try:
        n = exporter(dossier, args.sortie)
    except PermissionError:
        log("Base suivi magasins.xlsx est ouvert dans Excel, impossible de le lire. Fermez-le et relancez.")
        return 1
    log(f"{n} lignes exportees -> {args.sortie}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nInterrompu.")
        sys.exit(130)
